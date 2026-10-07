"""Teste de ponta a ponta offline: feeds e páginas simulados, redator mock, geração dos .docx."""
import io
import json
import shutil
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from scitekvision import coletor, extrator, pipeline, ranking, redator
from scitekvision.config import RAIZ, carregar

AGORA = datetime.now(timezone.utc)


def rss(itens):
    corpo = "".join(
        f"<item><title>{t}</title><link>{u}</link><description>{d}</description>"
        f"<pubDate>{(AGORA - timedelta(hours=h)).strftime('%a, %d %b %Y %H:%M:%S +0000')}</pubDate></item>"
        for t, u, d, h in itens)
    return f'<?xml version="1.0"?><rss version="2.0"><channel><title>x</title>{corpo}</channel></rss>'.encode()


def pagina(titulo):
    par = " ".join(f"Researchers reported finding number {i} about {titulo} in a peer reviewed study." for i in range(40))
    return (f'<html><head><title>{titulo}</title><meta property="og:image" content="https://img.test/{abs(hash(titulo))}.jpg">'
            f'<meta property="og:site_name" content="Site Teste"></head><body><article><h1>{titulo}</h1>'
            f"<p>{par}</p><p>{par}</p></article></body></html>")


def jpeg():
    b = io.BytesIO()
    Image.new("RGB", (800, 450), (20, 60, 140)).save(b, "JPEG")
    return b.getvalue()


class Resp:
    def __init__(self, content=b"", status=200):
        self.content, self.status_code = content, status
        self.text = content.decode("utf-8", "ignore") if isinstance(content, bytes) else content

    def raise_for_status(self):
        if self.status_code >= 400:
            import requests
            raise requests.HTTPError(str(self.status_code))

    def json(self):
        return json.loads(self.text)


@pytest.fixture
def cfg(tmp_path, monkeypatch):
    for d in ("config", "prompts"):
        shutil.copytree(RAIZ / d, tmp_path / d)
    c = carregar(tmp_path)
    c.geral["redacao"]["backends"] = ["mock"]

    def fake_get(url, **kw):
        if "img.test" in url:
            return Resp(jpeg())
        if url.startswith("https://noticia.test/"):
            return Resp(pagina(url.rsplit("/", 1)[-1]).encode())
        if "wikimedia" in url:
            return Resp(b'{"query": {"pages": {}}}')
        # qualquer feed configurado: 3 notícias por feed, com títulos ligados à editoria
        cat = next((k for k, v in c.categorias.items() if any(f["url"] == url for f in v.get("feeds", []))), None)
        if cat is None:
            return Resp(status=404)
        kw0 = c.categorias[cat]["palavras_chave"][0]
        return Resp(rss([(f"New {kw0} breakthrough {i} {cat}", f"https://noticia.test/{cat}-{i}-{abs(hash(url))}",
                          f"A study on {kw0}.", i * 3) for i in range(3)]))

    for mod in (coletor, extrator, redator):
        monkeypatch.setattr(mod.requests, "get", fake_get)
    return c


def test_edicao_completa(cfg):
    edicao = date.today() + timedelta(days=1)
    arquivos = pipeline.edicao_completa(cfg, edicao, ignorar_prazo=True)
    pasta = cfg.pasta_edicao(edicao)
    pautas = pipeline.carregar_pautas(cfg, edicao)
    obrigatorias = {k for k, v in cfg.categorias.items() if not v.get("opcional")}
    assert obrigatorias <= {p.categoria for p in pautas}
    assert all(p.status == "redigida" for p in pautas)
    assert all(p.imagem_arquivo and (pasta / p.imagem_arquivo).exists() for p in pautas)
    assert len({p.candidato.url for p in pautas}) == len(pautas)  # sem repetição entre editorias
    consolidado = arquivos[-1]
    assert consolidado.name.endswith("edicao-completa.docx")
    doc = Document(consolidado)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "RASCUNHO" in "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Fontes e créditos" in texto and len(doc.inline_shapes) == len(pautas)
    assert (pasta / "LEIA-ME-REVISAO.md").exists()


def test_prazo_bloqueia_redacao(cfg):
    ontem = date.today() - timedelta(days=1)
    coletor.executar(cfg, ontem)
    pipeline.selecionar(cfg, ontem)
    pipeline.extrair(cfg, ontem)
    pautas = pipeline.redigir(cfg, ontem)
    assert pautas and all(p.status == "fora_do_prazo" for p in pautas)
    arquivos = pipeline.montar(cfg, ontem)  # entrega mesmo assim, com pautas pendentes
    assert "PENDENTE" in "\n".join(p.text for p in Document(arquivos[0]).paragraphs)


def test_artigo_escrito_pelo_claude_code_e_incorporado(cfg):
    edicao = date.today() + timedelta(days=1)
    coletor.executar(cfg, edicao)
    pipeline.selecionar(cfg, edicao)
    pautas = pipeline.extrair(cfg, edicao)
    p = pautas[0]
    artigo = json.loads(redator._mock("", f"Título original: {p.candidato.titulo}", ""))
    (cfg.pasta_edicao(edicao) / "artigos").mkdir(parents=True, exist_ok=True)
    (cfg.pasta_edicao(edicao) / "artigos" / f"{p.categoria}.json").write_text(json.dumps(artigo), encoding="utf-8")
    recarregadas = pipeline.carregar_pautas(cfg, edicao)
    assert recarregadas[0].status == "redigida" and recarregadas[0].backend == "claude-code"
    assert (cfg.pasta_edicao(edicao) / "fontes" / f"{p.categoria}.md").exists()


def test_edicao_vira_no_prazo(cfg):
    tz = cfg.tz
    assert cfg.edicao_atual(datetime(2026, 10, 7, 9, 59, tzinfo=tz)) == date(2026, 10, 7)
    assert cfg.edicao_atual(datetime(2026, 10, 7, 22, 0, tzinfo=tz)) == date(2026, 10, 8)


def test_ranking_prefere_fonte_oficial_e_repercussao(cfg):
    from scitekvision.modelos import Candidato
    base = dict(resumo="", categorias=["astrofisica"], coletado_em=AGORA.isoformat(),
                publicado=(AGORA - timedelta(hours=2)).isoformat())
    pool = [
        Candidato(id="1", titulo="Webb telescope finds water on exoplanet", url="u1", fonte="NASA", peso_fonte=1.0, **base),
        Candidato(id="2", titulo="Webb finds water on distant exoplanet", url="u2", fonte="Space.com", peso_fonte=0.7, **base),
        Candidato(id="3", titulo="Best telescope deals today: 30% off", url="u3", fonte="Space.com", peso_fonte=0.7, **base),
    ]
    sel = ranking.selecionar(pool, cfg)
    assert [c.id for c in sel["astrofisica"]][0] == "1"
    assert sel["astrofisica"][-1].id == "3"


def test_detecta_copia():
    fonte = "the quick brown fox jumps over the lazy dog near the river bank today"
    art = {"lead": fonte, "corpo": []}
    assert redator.sobreposicao(art, fonte) > 0.9
