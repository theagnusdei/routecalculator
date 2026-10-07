"""Etapa 5 — montagem dos arquivos .docx para revisão humana.

Gera um .docx por artigo e um volume consolidado da edição, com capa, sumário,
status de revisão, imagem com crédito/licença, fontes e checklist do revisor.
"""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import Config, slug
from .modelos import Pauta

log = logging.getLogger(__name__)

AZUL = RGBColor(0x0B, 0x3D, 0x91)
CINZA = RGBColor(0x55, 0x55, 0x55)
VERMELHO = RGBColor(0xB0, 0x00, 0x20)

CHECKLIST = [
    "Fatos, números, nomes e datas conferidos com a fonte original",
    "Texto é original (sem trechos copiados) e está em bom português",
    "Título e linha fina são precisos, sem sensacionalismo",
    "Imagem: direito de uso confirmado (ou substituída por imagem livre) e crédito correto",
    "Fontes e links funcionando",
    "Alertas do redator automático resolvidos",
    "Aprovado para publicação",
]


def _estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)


def _sombrear(celula, cor_hex: str) -> None:
    tc_pr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    tc_pr.append(shd)


def _link(paragrafo, url: str, texto: str | None = None) -> None:
    rel = paragrafo.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                   is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    cor = OxmlElement("w:color")
    cor.set(qn("w:val"), "0B3D91")
    sub = OxmlElement("w:u")
    sub.set(qn("w:val"), "single")
    rpr.append(cor)
    rpr.append(sub)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto or url
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragrafo._p.append(h)


def _faixa_status(doc: Document, texto: str) -> None:
    tabela = doc.add_table(rows=1, cols=1)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = tabela.rows[0].cells[0]
    _sombrear(cel, "FFF4CE")
    run = cel.paragraphs[0].add_run(texto)
    run.bold = True
    run.font.color.rgb = VERMELHO
    cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _tabela_meta(doc: Document, linhas: list[tuple[str, str]]) -> None:
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in linhas:
        a, b = t.add_row().cells
        a.width, b.width = Cm(4.5), Cm(12)
        _sombrear(a, "EAF0FA")
        a.paragraphs[0].add_run(k).bold = True
        if v.startswith("http"):
            _link(b.paragraphs[0], v)
        else:
            b.paragraphs[0].add_run(v)


def escrever_artigo(doc: Document, pauta: Pauta, pasta: Path, cfg: Config, edicao: date) -> None:
    art, c = pauta.artigo or {}, pauta.candidato
    nome_cat = cfg.categorias[pauta.categoria]["nome"]

    _faixa_status(doc, "RASCUNHO — AGUARDANDO REVISÃO E APROVAÇÃO HUMANA")
    p = doc.add_paragraph()
    r = p.add_run(nome_cat.upper())
    r.bold, r.font.color.rgb, r.font.size = True, AZUL, Pt(10)

    if pauta.status != "redigida":
        doc.add_heading(f"[PENDENTE] {c.titulo}", level=1)
        doc.add_paragraph(f"O artigo desta editoria não foi redigido automaticamente "
                          f"(status: {pauta.status}; motivo: {pauta.erro or 'não informado'}). "
                          f"Pauta sugerida para redação manual:")
        _tabela_meta(doc, [("Matéria-fonte", c.url), ("Veículo", c.fonte), ("Resumo do feed", c.resumo or "—")])
        return

    doc.add_heading(art["titulo"], level=1)
    sub = doc.add_paragraph()
    rs = sub.add_run(art["subtitulo"])
    rs.italic, rs.font.size, rs.font.color.rgb = True, Pt(13), CINZA

    if pauta.imagem_arquivo and (pasta / pauta.imagem_arquivo).exists():
        doc.add_picture(str(pasta / pauta.imagem_arquivo), width=Cm(16))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        leg = doc.add_paragraph()
        leg.add_run(art.get("legenda_imagem", "")).italic = True
        cred = leg.add_run(f"  Imagem: {pauta.imagem_credito}. ")
        cred.font.size, cred.font.color.rgb = Pt(9), CINZA
        if pauta.imagem_pagina:
            _link(leg, pauta.imagem_pagina, "Fonte da imagem")
        lic = doc.add_paragraph()
        rl = lic.add_run(f"Licença/uso: {pauta.imagem_licenca}")
        rl.font.size, rl.font.color.rgb = Pt(8), CINZA
    else:
        aviso = doc.add_paragraph()
        aviso.add_run(f"[IMAGEM PENDENTE] Legenda sugerida: {art.get('legenda_imagem', '')}").bold = True

    lead = doc.add_paragraph()
    rl = lead.add_run(art["lead"])
    rl.bold, rl.font.size = True, Pt(12)

    for secao in art["corpo"]:
        if secao.get("intertitulo"):
            doc.add_heading(secao["intertitulo"], level=2)
        for par in secao.get("paragrafos", []):
            doc.add_paragraph(par).paragraph_format.space_after = Pt(8)

    doc.add_heading("Por que importa", level=2)
    doc.add_paragraph(art["por_que_importa"])
    doc.add_heading("Pontos-chave", level=2)
    for item in art["pontos_chave"]:
        doc.add_paragraph(item, style="List Bullet")
    if art.get("glossario"):
        doc.add_heading("Glossário", level=2)
        for g in art["glossario"]:
            gp = doc.add_paragraph(style="List Bullet")
            gp.add_run(f"{g.get('termo', '')}: ").bold = True
            gp.add_run(g.get("definicao", ""))

    doc.add_heading("Fontes e créditos", level=2)
    fp = doc.add_paragraph(f"Com informações de {c.fonte}: ")
    _link(fp, c.url, c.titulo)
    for u in c.cobertura[:5]:
        _link(doc.add_paragraph(style="List Bullet"), u)

    doc.add_heading("Para o revisor (remover antes de publicar)", level=2)
    _tabela_meta(doc, [
        ("Edição", edicao.strftime("%d/%m/%Y")),
        ("Editoria", nome_cat),
        ("Matéria-fonte", c.url),
        ("Autor original", pauta.autor_fonte or "não informado"),
        ("Publicada em", c.publicado or "não informado"),
        ("Relevância (score)", f"{c.score:.2f} — repercussão em {len(c.cobertura)} outra(s) fonte(s)"),
        ("Redigido por", f"IA ({pauta.backend}) — confiança declarada: {art.get('nivel_confianca', '?')}"),
        ("Palavras-chave (SEO)", ", ".join(art.get("palavras_chave", []))),
        ("Meta descrição", art.get("meta_descricao", "")),
    ])
    alertas = list(art.get("alertas_revisor", [])) + pauta.alertas_automaticos
    if alertas:
        doc.add_paragraph().add_run("Alertas:").bold = True
        for a in alertas:
            doc.add_paragraph(a, style="List Bullet")
    doc.add_paragraph().add_run("Checklist:").bold = True
    for item in CHECKLIST:
        doc.add_paragraph(f"☐ {item}")


def _novo_documento(titulo: str, cfg: Config) -> Document:
    doc = Document()
    _estilos(doc)
    doc.core_properties.title = titulo
    doc.core_properties.author = f"{cfg.geral['saida']['nome_publicacao']} — agente editorial"
    doc.core_properties.comments = "Rascunho gerado automaticamente; requer revisão humana."
    return doc


def nome_arquivo(i: int, pauta: Pauta) -> str:
    titulo = (pauta.artigo or {}).get("titulo") or pauta.candidato.titulo
    return f"{i:02d}-{pauta.categoria}-{slug(titulo, 50)}.docx"


def montar(pautas: list[Pauta], pasta: Path, cfg: Config, edicao: date) -> list[Path]:
    pub = cfg.geral["saida"]["nome_publicacao"]
    (pasta / "docx").mkdir(parents=True, exist_ok=True)
    gerados = []
    for i, pauta in enumerate(pautas, 1):
        doc = _novo_documento((pauta.artigo or {}).get("titulo", pauta.candidato.titulo), cfg)
        escrever_artigo(doc, pauta, pasta, cfg, edicao)
        destino = pasta / "docx" / nome_arquivo(i, pauta)
        doc.save(destino)
        gerados.append(destino)

    # volume consolidado
    doc = _novo_documento(f"{pub} — edição de {edicao:%d/%m/%Y}", cfg)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = t.add_run(pub)
    rt.bold, rt.font.size, rt.font.color.rgb = True, Pt(36), AZUL
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run(f"Edição de {edicao:%d/%m/%Y} — rascunhos para revisão").font.size = Pt(14)
    redigidos = sum(p.status == "redigida" for p in pautas)
    doc.add_paragraph(f"{redigidos} de {len(pautas)} editorias com artigo redigido.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Sumário", level=1)
    tab = doc.add_table(rows=1, cols=4)
    tab.style = "Table Grid"
    for cel, txt in zip(tab.rows[0].cells, ["#", "Editoria", "Título", "Status"]):
        cel.paragraphs[0].add_run(txt).bold = True
        _sombrear(cel, "EAF0FA")
    for i, p in enumerate(pautas, 1):
        row = tab.add_row().cells
        row[0].text = str(i)
        row[1].text = cfg.categorias[p.categoria]["nome"]
        row[2].text = (p.artigo or {}).get("titulo") or p.candidato.titulo
        row[3].text = "Rascunho pronto" if p.status == "redigida" else f"PENDENTE ({p.status})"
    for linha in tab.rows:
        for cel, largura in zip(linha.cells, (1.0, 3.8, 9.2, 3.0)):
            cel.width = Cm(largura)

    for p in pautas:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        escrever_artigo(doc, p, pasta, cfg, edicao)
    consolidado = pasta / f"{pub}-{edicao.isoformat()}-edicao-completa.docx"
    doc.save(consolidado)
    gerados.append(consolidado)
    log.info("%d arquivos .docx gerados em %s", len(gerados), pasta)
    return gerados
